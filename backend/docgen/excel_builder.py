import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


def build_spreadsheet(structure: dict) -> bytes:
    """Builds a real .xlsx from a {"sheet_title": str, "headers": [str],
    "rows": [[str|number]], "chart_title": str | None} structure (the
    shape agents.tasks._run_excel_generation_step parses from the
    preceding step's raw text). Malformed/missing chart data is skipped,
    not fatal — the sheet still ships without a chart rather than failing
    the run."""
    workbook = Workbook()
    sheet = workbook.active
    # Excel sheet names are capped at 31 characters.
    sheet.title = str(structure.get("sheet_title", "Sheet1"))[:31] or "Sheet1"

    headers = structure.get("headers") or []
    if headers:
        sheet.append([str(header) for header in headers])

    rows = structure.get("rows") or []
    for row in rows:
        sheet.append([_coerce_cell(value) for value in row])

    _add_chart_if_valid(sheet, structure.get("chart_title"), headers, rows)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _coerce_cell(value):
    if isinstance(value, (int, float)):
        return value
    try:
        return float(value)
    except (TypeError, ValueError):
        return str(value)


def _add_chart_if_valid(sheet, chart_title, headers: list, rows: list) -> None:
    if not chart_title or len(headers) < 2 or not rows:
        return
    try:
        chart = BarChart()
        chart.title = str(chart_title)
        data = Reference(
            sheet,
            min_col=2,
            max_col=len(headers),
            min_row=1,
            max_row=1 + len(rows),
        )
        categories = Reference(
            sheet, min_col=1, min_row=2, max_row=1 + len(rows)
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        sheet.add_chart(chart, f"A{len(rows) + 3}")
    except (KeyError, TypeError, ValueError, AttributeError):
        pass
